from typing import Any, Text, Dict, List, Optional
from rasa_sdk import Action, Tracker
from rasa_sdk.executor import CollectingDispatcher
from rasa_sdk.events import SlotSet, SessionStarted, ActionExecuted, FollowupAction
import os
import re
import json
from datetime import datetime
from pathlib import Path
import joblib
import nltk
from nltk.corpus import stopwords
from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from dotenv import load_dotenv
from langchain_openai import ChatOpenAI
from langchain_core.prompts import PromptTemplate
from langchain_core.runnables import RunnablePassthrough, RunnableParallel
 

# Asegurar que las stopwords estén descargadas
nltk.download('stopwords')
stop_words = set(stopwords.words('spanish'))

# Cargar modelo y vectorizador
modelo = joblib.load("logistic_regression.joblib")
vectorizador = joblib.load("vectorizer_tfidf.joblib")

# Cargar .env
dotenv_path = os.path.join(os.path.dirname(__file__), "../../.env")
load_dotenv(dotenv_path=dotenv_path)

class ActionResumenYClasificacion(Action):
    def name(self) -> str:
        return "action_resumen_y_clasificacion"

    def run(self, dispatcher: CollectingDispatcher,
            tracker: Tracker,
            domain: Dict[str, Any]) -> List[Dict[str, Any]]:

        modo_simulado = False # TRUE = Simulación (no gasta tokens de openai), poner en False solo si se quiere probar el modelo con IA generativa
        
        resumen_es = None #Empezamos con el resumen vacío por si hay intentos previos guardados en el chat
        chat_str = None # Inicializamos el string del chat vacío también por si hay intentos previos guardados en el chat

        # --- Extraer JSON desde utter_inform_perfil_victima1 ---
        extraer = False
        incluir_siguiente_bot_text = False
        chat = []
        ultimo_utter = None

        for event in tracker.events:
            if event.get("event") == "action": 
                nombre_accion = event.get("name")

                if nombre_accion == "utter_inform_perfil_victima1":
                    chat = []
                    extraer = True
                    incluir_siguiente_bot_text = True

                if nombre_accion.startswith("utter_"):
                    ultimo_utter = nombre_accion                        
                continue

            if extraer:
                if incluir_siguiente_bot_text and event.get("event") == "bot" and event.get("text"):
                    chat.append({
                        "role": "chatbot",
                        "message": event.get("text"),
                        "utter_action": ultimo_utter
                    })
                    incluir_siguiente_bot_text = False
                    continue

                if event.get("event") == "user" and event.get("text"):
                    chat.append({
                        "role": "usuario",
                        "message": event.get("text")
                    })

                elif event.get("event") == "bot" and event.get("text"):
                    chat.append({
                        "role": "chatbot",
                        "message": event.get("text"),
                        "utter_action": ultimo_utter
                    })

        chat_str = json.dumps({"chat": chat}, ensure_ascii=False)

        if not chat:
            print(f"Ocurrió un error interno: No se pudo extraer el chat.") #PARA DEBUG
            return [SlotSet("chat", chat_str), SlotSet("resumen_es", resumen_es)] 
        
        print(f"DEBUG: JSON extraído:\n```\n{json.dumps({'chat': chat}, ensure_ascii=False, indent=2)}") #PARA DEBUG
        
        # Validar el relato para evitar testimonios vacíos
        def extraer_mensaje_relato(chat: List[Dict[str, str]]) -> Optional[str]:
            siguiente_es_relato = False
            for mensaje in chat:
                if mensaje["role"] == "chatbot" and mensaje.get("utter_action", "").startswith("utter_invita_relato_hechos"):
                    siguiente_es_relato = True
                    continue
                if siguiente_es_relato and mensaje["role"] == "usuario":
                    return mensaje["message"].strip()
            return None

        def relato_valido(texto: str) -> bool:
            return bool(texto and len(texto.split()) >= 3)

        mensaje_relato = extraer_mensaje_relato(chat)
        if not modo_simulado and (not mensaje_relato or not relato_valido(mensaje_relato)):
            print("No hay información suficiente para generar un resumen útil.")  # PARA DEBUG
            return [SlotSet("chat", chat_str), SlotSet("resumen_es", None)]

        api_key = os.getenv("OPENAI_API_KEY")
        if not api_key:
            print("Ocurrió un error interno: MissingAPIKey")
            return [SlotSet("chat", chat_str), SlotSet("resumen_es", resumen_es)]

        # --- Guardrail para conversaciones troll o incompletas ---

        def es_conversacion_valida(chat_str: str, api_key: str, modo_simulado: bool) -> bool:
            if modo_simulado:   #COMENTAR SI SE QUIERE PROBAR EL GUARDRAIL DURANTE EL MODO SIMULADO
                return True

            prompt_text = """
            A continuación se proporciona una conversación entre un usuario y un asistente diseñado para ayudar a generar informes sobre posibles delitos de odio.

            Tu tarea es analizar la conversación y determinar si el usuario está compartiendo un testimonio legítimo o si está haciendo un uso inadecuado del sistema (por ejemplo, escribiendo bromas, respuestas incoherentes, lenguaje burlón, o sin aportar información relevante). 
            Si no estás seguro, considera que la conversación es válida.
            
            Devuelve únicamente una palabra (**siempre en español**):
            - válido → si la conversación parece auténtica.
            - troll → si parece una broma, un engaño o un contenido inapropiado.

            ### Conversación:
            {chat}
            """

            prompt = PromptTemplate.from_template(prompt_text)
            llm = ChatOpenAI(model="gpt-3.5-turbo", temperature=0, api_key=api_key)

            chain = (RunnableParallel(chat=RunnablePassthrough()) | prompt | llm)
            respuesta = chain.invoke(chat_str).content.strip().lower()

            return respuesta == "válido"
        
        if not es_conversacion_valida(chat_str, api_key, modo_simulado):    
            print("Conversación descartada por contenido inadecuado.")  # PARA DEBUG
            return [SlotSet("chat", chat_str), SlotSet("resumen_es", resumen_es)]
        
        # --- Generar resumen con LLM ---

        if modo_simulado:

            # Simulamos un resumen de prueba para no gastar tokens
            resumen_es = """
            Informe testimonial

            1. Datos del usuario:
            - Edad: 30 años
            - Género: No se especifica.
            - Nacionalidad: Española.
            - Pertenencia a grupo o comunidad relevante: El usuario se identifica como persona trans y forma parte de un colectivo LGTBIQ+.

            2. Ubicación de los hechos:
            - Fecha del incidente: Aproximadamente el 10 de abril de 2025.
            - Medio en que se sucedieron los hechos: Redes sociales.
            - Lugar del incidente: Plataforma Instagram, en los comentarios de una publicación personal del usuario.

            3. Descripción de los hechos:
            El usuario publicó una fotografía con un mensaje sobre visibilidad trans. Horas después, comenzó a recibir comentarios ofensivos y transfóbicos por parte de varias cuentas anónimas. Entre los mensajes recibidos se encontraban expresiones como: "esto no es normal", "deberías meterte en un psiquiátrico" y "ojalá desaparezcáis".

            Una de las cuentas más activas utilizaba el nombre de usuario "@limpieza_social2025" y compartía contenido de ideología extremista. Esta cuenta llegó a escribir: “hay que hacer limpieza, como en los viejos tiempos”.

            El usuario bloqueó varias cuentas y denunció los comentarios a la plataforma, aunque hasta el momento no ha recibido respuesta oficial. 
            """

        else:
            
            prompt_template =prompt_template = """
            ### Rol ###
            Eres un asistente especializado en la extracción de información y redacción de informes a partir de testimonios de personas que podrían haber presenciado o sufrido situaciones discriminatorias o violentas.

            ### Objetivo ###
            Tu tarea es extraer y organizar exclusivamente la información que aparece de forma explícita en el testimonio. Luego, redacta un informe claro y estructurado, sin añadir interpretaciones, suposiciones ni contenido adicional. El testimonio se proporciona en formato JSON como una conversación entre un usuario y un chatbot.

            ---

            ### Instrucciones generales ###
            - Se te proporcionará un chat en formato JSON.
            - El informe debe basarse **únicamente** en la información que aparece en el JSON.
            - ❗ **No infieras, completes ni inventes datos que no estén claramente indicados.**
            - Si una pregunta no puede responderse con la información dada, **omite esa parte sin hacer mención de su ausencia**.
            - Si se proporcionan palabras o frases que fueron pronunciadas durante los hechos, **reprodúcelas como citas textuales entre comillas**.
            - Si el usuario no declara explícitamente su género **completa el campo "Género" con las palabras "No se especifica"**.
            - Si durante la conversación se le ha ofrecido información o apoyo al usuario, **omite esa parte sin hacer mención de su ausencia**.
            - Redacta en un lenguaje **formal pero accesible**, claro y empático.
            - Organiza la información de forma estructurada según los apartados que siguen.

            ---

            ### Esquema del informe ###

            #### 1. Datos del usuario
            - Edad.
            - Género.
            - Nacionalidad.
            - Pertenencia a algún grupo o comunidad relevante (etnia, religión, orientación sexual, identidad de género, discapacidad, etc.).

            #### 2. Ubicación de los hechos
            - Fecha y hora (si se indican) y su posible relevancia.
            - Medio en que sucedieron los hechos (presencial, redes sociales, etc.).
            - Lugar donde ocurrieron los hechos y su posible relevancia (cercanías de un lugar o acto simbólico, por ejemplo).

            #### 3. Descripción de los hechos teniendo en cuenta:
            - Relato detallado del testimonio del usuario.
            - Palabras, frases o insultos pronunciados (en citas textuales).
            - Si el usuario conocía a la(s) persona(s) implicada(s).
            - Mención de posibles antecedentes legales de los implicados (si se menciona).
            - Descripción de los agresores (física, simbología, vestimenta, etc., especialmente si hace referencia a grupos extremistas).
            - Si se llamó a las autoridades y cómo intervinieron (si se indica).
            - Presencia de otros testigos y la información disponible sobre ellos.
            - Si el usuario ha denunciado los hechos.

            ---

            ### Recuerda ###
            - No agregues información que no esté explícitamente en el JSON.
            - No completes huecos ni reformules hechos con lenguaje no presente en el testimonio.
            - Cíñete estrictamente a lo expresado por el usuario y a los apartados proporcionados.
            - No menciones si en el chat se ha ofrecido información o apoyo al usuario.
            - No indiques el género del usuario si no lo ha declarado explícitamente.

            ### Testimonio proporcionado (formato JSON) ###
            {json}

            ### Formato de ejemplo (no usar este contenido, solo imitar el formato) ###

            Informe testimonial

            1. Datos del usuario:
            - Edad: 26 años
            - Género: No se indica.
            - Nacionalidad: No se indica.
            - Pertenencia a grupo o comunidad relevante: Es homosexual y miembro de un colectivo LGTBIQ+

            2. Ubicación de los hechos:
            - Fecha del incidente: 18 de marzo de 2024.
            - Hora del incidente: 23:30 horas.
            - Medio en que se sucedieron los hechos: De forma presencial.
            - Lugar del incidente: Calle céntrica, en las proximidades de un bar LGTBIQ+ llamado "El Paso".

            3. Descripción de los hechos:
            El usuario, acompañado de un grupo de amigos, salió de un bar LGTBIQ+ cuando un grupo de tres individuos desconocidos comenzó a seguirles. Uno de ellos profirió insultos homofóbicos en voz alta, expresando: “Mira a estos maricones, seguro que van buscando problemas”.
            Ante esta situación, el usuario y su grupo intentaron ignorar a los agresores. Sin embargo, los individuos se acercaron más, intensificando la agresión verbal y física. Uno de los agresores empujó al usuario y le dirigió las palabras: “Deberíais desaparecer de aquí, dais asco”. Posteriormente, cuando el usuario intentó alejarse, uno de los agresores le propinó un golpe en el brazo y lo hizo caer al suelo.
            - Descripción de los agresores:
            Tres individuos de entre 20 y 25 años de edad. Uno de ellos vestía una chaqueta negra con un símbolo asociado presuntamente a un grupo de ultraderecha. Otro portaba una gorra roja.
            El tercer agresor llevaba una sudadera gris. El usuario no los había visto con anterioridad.
            - Testigos del incidente:
            Amigos del usuario presentes en el lugar de los hechos.
            Una pareja que transitaba por la calle en el momento del ataque.
            Un camarero de un bar cercano, quien observó lo ocurrido y procedió a llamar a la policía.
            - Intervención de las autoridades:
            A pesar de que la policía llegó al lugar tras la llamada del camarero, los agresores ya habían huido corriendo. Posteriormente, el usuario se dirigió a la comisaría correspondiente y formalizó la denuncia ante las autoridades pertinentes.
            """

            try:
                prompt = PromptTemplate(template=prompt_template, input_variables=["json"])
                llm = ChatOpenAI(model="gpt-3.5-turbo", temperature=0, api_key=api_key)
                assistant_chain = (RunnableParallel(json=RunnablePassthrough()) | prompt | llm)
                result = assistant_chain.invoke(chat_str)
                resumen_es = result.content
                if not resumen_es:
                    dispatcher.utter_message(text="Ocurrió un error interno: No se pudo generar el resumen.")
                    return [SlotSet("chat", chat_str), SlotSet("resumen_es", resumen_es)]

            except Exception as e:
                dispatcher.utter_message(text=f"Ocurrió un error interno al generar el resumen: {type(e).__name__}")
                return [SlotSet("chat", chat_str), SlotSet("resumen_es", resumen_es)]


        # --- Clasificar discurso de odio ---
          
        # --- Tomar narración -----
        match = re.search(r"(?s)3\. Descripción de los hechos:\s*(.+?)(?:\n\d\.|\Z)", resumen_es)
        descripcion_hechos = match.group(1).strip() if match else resumen_es


        # --- Clasificación---
        texto_limpio = re.sub(r"[^a-záéíóúñü\s]", "", descripcion_hechos.lower())
        texto_limpio = " ".join([p for p in texto_limpio.split() if p not in stop_words])
        X = vectorizador.transform([texto_limpio])
        probabilidades = modelo.predict_proba(X)[0]
        clase = int(probabilidades.argmax())
        confianza = round(probabilidades[clase] * 100, 2)

        clases = {
            0: "El sistema no detecta indicios claros que permitan identificar los hechos como un posible delito de odio.",
            1: "El sistema identifica elementos que podrían constituir un delito de odio.",
            2: "El sistema identifica elementos que son problemáticos o están en el límite (borderline)."
        }

        dispatcher.utter_message(text=clases[clase] + f" (Confianza: {confianza}%)")

        # Guardar slots
        return [
            SlotSet("chat", chat_str),
            SlotSet("resumen_es", resumen_es)
        ]


class ActionGenerarInforme(Action):
    def name(self) -> Text:
        return "action_generar_informe"

    def run(self, dispatcher: CollectingDispatcher,
            tracker: Tracker,
            domain: Dict[Text, Any]) -> List[Dict[Text, Any]]:

        # Obtener el resumen desde el slot
        resumen_es = tracker.get_slot("resumen_es")
        if not resumen_es:
            dispatcher.utter_message(text="Error de validación: La información disponible es insuficiente o inadecuada para generar un documento.")
            return []

        # Crear el documento Word
        doc = Document()
        now = datetime.now()
        fecha_str = now.strftime("%d/%m/%Y a las %H:%M")
        file_name = f"documento_manifestacion_personal_{now.strftime('%Y%m%d_%H%M%S')}.docx"

        # Agregar título centrado
        title = doc.add_heading("Documento de manifestación personal", 0)
        title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

        # Agregar fecha de generación
        intro = doc.add_paragraph()
        intro.add_run(f"Este informe ha sido generado automáticamente por un sistema de inteligencia artificial el día {fecha_str}. Su función es meramente informativa y no tiene ninguna validez legal.").italic = True
        doc.add_paragraph("")  # Línea en blanco

        # Procesar el resumen línea por línea
        lines = resumen_es.strip().split('\n')
        if lines[0].startswith("Informe"):
            lines = lines[1:]

        for line in lines:
            line = line.strip()
            if line.startswith(("1.", "2.", "3.", "Datos del usuario", "Ubicación de los hechos", "Descripción de los hechos")):
                doc.add_paragraph(line, style="Heading 2")
            elif line.startswith("- "):
                doc.add_paragraph(line.lstrip("- ").strip(), style="List Bullet")
            elif line:
                doc.add_paragraph(line)

        # Guardar en carpeta accesible
        docs_dir = Path(__file__).resolve().parents[2] / "backend" / "docs"
        os.makedirs(docs_dir, exist_ok=True)
        file_path = docs_dir / file_name
        doc.save(file_path)

        # Crear URL para descarga
        download_url = f"http://localhost:8000/docs/{file_name}"
        mensaje_base = (
            'Tu documento está listo. '
            f'<a href="{download_url}" target="_blank">Haz clic aquí</a> para descargarlo.'
        )
        dispatcher.utter_message(text=f"##INFORME## {mensaje_base} ||| {file_name}")
        
        return []


    
class ActionSessionStart(Action):
    def name(self) -> Text:
        return "accion_inicio_sesion"

    async def run(
        self, dispatcher: CollectingDispatcher, tracker: Tracker, domain: Dict[Text, Any]
    ) -> List[Dict[Text, Any]]:
        buttons = [
            {"title": "Quiero entender mejor qué se considera un delito de odio.", "payload": "/delito_odio"},
            {"title": "Quiero saber dónde y cómo puedo denunciar.", "payload": "/recursos_denuncia"},
            {"title": "Quiero contar lo que me pasó y recibir ayuda personalizada.", "payload": "/testimonio"},
        ]

        dispatcher.utter_message(
            text="¡Hola! Soy DEL.IA, un asistente virtual centrado en apoyarte si has vivido un acto de odio. Cuéntame, ¿con qué te gustaría que te ayude?",
            buttons=buttons
        )
        return []

class ActionDelitoOdio(Action):
    def name(self) -> Text:
        return "action_delito_odio"

    def intent_ya_usado(self, tracker: Tracker, intent_objetivo: Text) -> bool:
        for event in tracker.events:
            if event.get("event") == "user":
                intent = event.get("parse_data", {}).get("intent", {}).get("name")
                if intent == intent_objetivo:
                    return True
        return False

    async def run(self, dispatcher: CollectingDispatcher,
                  tracker: Tracker, domain: Dict[Text, Any]) -> List[Dict[Text, Any]]:

        texto = (
            "Es estupendo que quieras aprender más sobre este tema.\n"
            "Un delito de odio se produce cuando alguien ataca, insulta o trata mal a otra persona solo por ser quien es. "
            "Normalmente, se utiliza como insulto el color de piel, las ideas, la religión, el país de procedencia, la orientación sexual, la situación económica, el sexo (hombre o mujer) o la existencia de una enfermedad o una discapacidad.\n"
            "En cuanto al agresor, no se trata solo de lo que hace, sino también de lo que dice o comparte y si eso provoca odio o violencia hacia otras personas.\n"
            "Si te ha pasado algo así o lo has visto, no dudes en pedir ayuda. Recuerda que estoy aquí para:"
        )

        botones = []
        if not self.intent_ya_usado(tracker, "recursos_denuncia"):
            botones.append({"title": "Contarte cómo y dónde denunciar.", "payload": "/recursos_denuncia"})
        if not self.intent_ya_usado(tracker, "testimonio"):
            botones.append({"title": "Acompañarte en función de tu historia, si quieres contarme lo que ha sucedido.", "payload": "/testimonio"})

        dispatcher.utter_message(text=texto, buttons=botones)
        return []

class ActionRecursosDenuncia(Action):
    def name(self) -> Text:
        return "action_recursos_denuncia"

    def intent_ya_usado(self, tracker: Tracker, intent_objetivo: Text) -> bool:
        for event in tracker.events:
            if event.get("event") == "user":
                intent = event.get("parse_data", {}).get("intent", {}).get("name")
                if intent == intent_objetivo:
                    return True
        return False

    async def run(self, dispatcher: CollectingDispatcher,
                  tracker: Tracker, domain: Dict[Text, Any]) -> List[Dict[Text, Any]]:

        texto = (
            "Si has sido testigo o te has enfrentado a una situación que crees que puede ser un delito de odio, puedes denunciar en tu comisaría más cercana o ponerte en contacto con Diaconía \n"
            "a través del número +34 677 614 068 o mandando un correo electrónico a sinetiquetassinodio@diaconia.es, dónde abogados y psicólogos te apoyaran y explicarán los pasos a seguir.\n "
            "También puedes contar conmigo para:"
        )

        botones = []
        if not self.intent_ya_usado(tracker, "delito_odio"):
            botones.append({"title": "Ayudarte a entender qué es un delito de odio.", "payload": "/delito_odio"})
        if not self.intent_ya_usado(tracker, "testimonio"):
            botones.append({"title": "Hablar de lo que ha pasado, te escucharé con atención y podré ayudarte más en función de tu historia.", "payload": "/testimonio"})

        dispatcher.utter_message(text=texto, buttons=botones)
        return [] 

class ActionNoResponde(Action):
    def name(self) -> Text:
        return "action_no_responde"
    
    def run(self, dispatcher: CollectingDispatcher,
            tracker: Tracker,
            domain: Dict[Text, Any]) -> List[Dict[Text, Any]]:
        # Utters que no deben contarse como la última pregunta real
        excluded_utterances = {"utter_proposito_cuestionario"}

        last_bot_utterance = None
        for event in reversed(tracker.events):
            if event.get("event") == "bot":
                utter_action = event.get("metadata", {}).get("utter_action")
                if utter_action and utter_action not in excluded_utterances:
                    last_bot_utterance = utter_action
                    break

        num_agresores = tracker.get_slot("num_agresores") or "uno"

        # Flujos de agresores por separado
        singular_flow = {
            "utter_relacion_agresor": "utter_preguntar_detalles_iniciales_unico",
            "utter_preguntar_detalles_iniciales_unico": "utter_preguntar_descripcion_agresor",
            "utter_preguntar_descripcion_agresor": "utter_preguntar_caracteristicas_distintivas_grupo_odio",
            "utter_preguntar_caracteristicas_distintivas_grupo_odio": "utter_testigos"
        }

        plural_flow = {
            "utter_relacion_agresor_plural": "utter_preguntar_detalles_iniciales_plural",
            "utter_preguntar_detalles_iniciales_plural": "utter_preguntar_descripcion_agresor_plural",
            "utter_preguntar_descripcion_agresor_plural": "utter_preguntar_caracteristicas_distintivas_grupo_odio_plural",
            "utter_preguntar_caracteristicas_distintivas_grupo_odio_plural": "utter_testigos"
        }

        if num_agresores == "varios" and last_bot_utterance in plural_flow:
            dispatcher.utter_message(text="Entiendo, no pasa nada, seguimos con el cuestionario.")
            return [ActionExecuted("action_listen"), FollowupAction(next_actions[last_bot_utterance])]
        elif num_agresores == "uno" and last_bot_utterance in singular_flow:
            dispatcher.utter_message(text="Entiendo, no pasa nada, seguimos con el cuestionario.")
            return [ActionExecuted("action_listen"), FollowupAction(next_actions[last_bot_utterance])]
        
        # Define the flow transitions based on the last question
        next_actions = {
            # Initial flow
            "utter_testimonio": "utter_pedir_confirmacion",
            "utter_pedir_confirmacion": "utter_despedida_concienciadora",
            "utter_inform_perfil_victima1": "utter_inform_perfil_victima2",
            "utter_inform_perfil_victima2": "utter_ubicar_medio",

            
            # Digital medium flow
            "utter_ubicar_medio": "utter_plataforma_redes",
            "utter_plataforma_redes": "utter_fecha_redes",
            "utter_fecha_redes": "utter_num_agresores_redes",
            "utter_num_agresores_redes": (
                "utter_pregunta_cuenta_agresores_vinculacion_redes"
                if num_agresores == "varios"
                else "utter_pregunta_cuenta_agresor_vinculacion_redes"
            ),
            "utter_pregunta_cuenta_agresor_vinculacion_redes": "utter_invita_relato_hechos",
            "utter_pregunta_cuenta_agresores_vinculacion_redes": "utter_invita_relato_hechos",
            "utter_invita_relato_hechos": (
                "utter_invita_relato_adicional_motivacion_plural"
                if num_agresores == "varios"
                else "utter_invita_relato_adicional_motivacion"
            ),
            
            # Physical medium flow
            "utter_pregunta_fecha_vida_real": "utter_pregunta_localizacion",
            "utter_pregunta_localizacion": "utter_num_agresores_vida_real",
            "utter_num_agresores_vida_real": "utter_invita_relato_hechos",
            
            # Common flows
            "utter_invita_relato_hechos": "utter_invita_relato_adicional_motivacion",
            "utter_invita_relato_adicional_motivacion": "utter_denuncia_hechos",
            "utter_invita_relato_adicional_motivacion_plural": "utter_denuncia_hechos",
            
            # Aggressor description flows
            "utter_relacion_agresor": "utter_preguntar_detalles_iniciales_unico",
            "utter_preguntar_detalles_iniciales_unico": "utter_preguntar_descripcion_agresor",
            "utter_preguntar_descripcion_agresor": "utter_preguntar_caracteristicas_distintivas_grupo_odio",
            "utter_preguntar_caracteristicas_distintivas_grupo_odio": "utter_testigos",
            "utter_relacion_agresor_plural": "utter_preguntar_detalles_iniciales_plural",
            "utter_preguntar_detalles_iniciales_plural": "utter_preguntar_descripcion_agresor_plural",
            "utter_preguntar_descripcion_agresor_plural": "utter_preguntar_caracteristicas_distintivas_grupo_odio_plural",
            "utter_preguntar_caracteristicas_distintivas_grupo_odio_plural": "utter_testigos",
            
            # Witness flow
            "utter_testigos": "utter_denuncia_hechos",
            "utter_relacion_testigo": "utter_info_testigos",
            "utter_info_testigos": "utter_denuncia_hechos",
            
            # Report flow
            "utter_recordatorio_redactar_informe": "utter_informe_no",
        }
        
        if last_bot_utterance in next_actions:
            dispatcher.utter_message(text="Entiendo, no pasa nada, seguimos con el cuestionario.")
            return [ActionExecuted("action_listen"), FollowupAction(next_actions[last_bot_utterance])]

 
class ActionSetRolAgresor(Action):
    def name(self):
        return "action_set_rol_agresor"
 
    def run(self, dispatcher, tracker, domain):
        return [SlotSet("rol", "agresor")]
 
class ActionSetRolTestigo(Action):
    def name(self):
        return "action_set_rol_testigo"
 
    def run(self, dispatcher, tracker, domain):
        return [SlotSet("rol", "testigo")]

class ActionSetTipoRedes(Action):
   def name(self):
       return "action_set_tipo_redes"
 
   def run(self, dispatcher, tracker, domain):
       return [SlotSet("tipo_fecha", "redes")]
 
class ActionSetTipoPersona(Action):
   def name(self):
       return "action_set_tipo_persona"
 
   def run(self, dispatcher, tracker, domain):
       return [SlotSet("tipo_fecha", "persona")]

from typing import Any, Text, Dict, List
from rasa_sdk import Action, Tracker
from rasa_sdk.executor import CollectingDispatcher
from rasa_sdk.events import UserUtteranceReverted
 
class ActionFallbackConReintento(Action):
 
    def name(self) -> Text:
        return "action_fallback_con_reintento"
 
    def run(self, dispatcher: CollectingDispatcher,
            tracker: Tracker,
            domain: Dict[Text, Any]) -> List[Dict[Text, Any]]:
 
        # Mensaje de disculpa
        dispatcher.utter_message(text="Lo siento, pero no te puedo ayudar con eso.")
 
        # Buscar la última acción del bot antes del fallback
        last_bot_utter = None
        for event in reversed(tracker.events):
            if event.get("event") == "action" and event.get("name", "").startswith("utter_"):
                last_bot_utter = event["name"]
                break
 
        if last_bot_utter:
            # Repetir el último mensaje del bot
            dispatcher.utter_message(response=last_bot_utter)
 
        # Opciónally revert user message to not affect the flow
        return [UserUtteranceReverted()]
